from __future__ import annotations

import copy
import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\CaNi517\OneDrive - HP Inc\Desktop\Portafolio de titulo\Fase 2\Haircut_Home_Studio\Fase 2\Evidencias Proyecto\HaircutHomeStudio")
REFERENCE = Path(r"C:\Users\CaNi517\Downloads\Documento Arquitectura Sistema(DAS).docx")
OUTPUT_DIR = ROOT / "docs" / "DAS"
OUTPUT = OUTPUT_DIR / "Documento Arquitectura Sistema Haircut Home Studio.docx"
DIAGRAMS = ROOT / "docs" / "diagramas"

EXPECTED_HASH = "619C683E5689B612B0D5043A5D3A625B16BBDA43FFEA701CA9A2D902273F8D00"
BLUE = RGBColor(0x00, 0x00, 0x80)
MEDIUM_BLUE = RGBColor(0x1F, 0x49, 0x7D)
BLACK = RGBColor(0x00, 0x00, 0x00)
LIGHT_BORDER = "D9D9D9"


def set_font(run, size=9.5, bold=False, color=BLACK, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def replace_paragraph(paragraph, text):
    target = next((r for r in paragraph.runs if r.text), paragraph.runs[0] if paragraph.runs else paragraph.add_run())
    for run in paragraph.runs:
        run.text = ""
    target.text = text


def set_cell_text(cell, text, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    run = p.add_run(text)
    set_font(run, size={1: 14, 2: 12, 3: 10}[level], bold=(level < 3), color=BLUE)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(5)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        marker = p.add_run("• ")
        set_font(marker)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            set_font(r, bold=True)
            r = p.add_run(rest)
            set_font(r)
        else:
            r = p.add_run(item)
            set_font(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5)
        shade_cell(table.rows[0].cells[i], "1F4E78")
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 22 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], value, align=align, size=8.5)
            if ri % 2 == 1:
                shade_cell(cells[i], "F2F6FA")
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, title, image_path, width=6.0, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    set_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run()
    r.add_picture(str(image_path), width=Inches(width))


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "La tabla de contenidos se actualizará al abrir el documento."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def remove_from_paragraph(doc, paragraph):
    body = doc._element.body
    start = paragraph._p
    found = False
    for child in list(body):
        if child is start:
            found = True
        if found and child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    for level, size in ((1, 14), (2, 12), (3, 10)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = level < 3
        style.font.color.rgb = BLUE


def build():
    digest = hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper()
    if digest != EXPECTED_HASH:
        raise RuntimeError("La plantilla cambió; se requiere una nueva inspección.")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    configure_styles(doc)
    doc.core_properties.title = "Documento de Arquitectura de Sistema Haircut Home Studio"
    doc.core_properties.subject = "Arquitectura de software 4+1"
    doc.core_properties.keywords = "Haircut Home Studio, DAS, arquitectura, Laravel, 4+1"

    # Portada: conservar composición e identidad visual.
    cover_name = next(p for p in doc.paragraphs if "Código Trauma" in p.text)
    cover_title = next(p for p in doc.paragraphs if "Software Architecture Document" in p.text)
    cover_version = next(p for p in doc.paragraphs if "Versión 1.0" in p.text)
    replace_paragraph(cover_name, "[Haircut Home Studio]")
    replace_paragraph(cover_title, "(DAS) Documento de Arquitectura de Sistema")
    replace_paragraph(cover_version, "Versión 1.0")

    # Control documental.
    set_cell_text(doc.tables[0].cell(0, 1), "DAS01")
    set_cell_text(doc.tables[0].cell(1, 1), "Haircut Home Studio")
    set_cell_text(doc.tables[0].cell(2, 1), "1.0")
    set_cell_text(doc.tables[1].cell(0, 1), "Equipo de Proyecto")
    set_cell_text(doc.tables[1].cell(1, 1), "27-09-2026")
    set_cell_text(doc.tables[1].cell(2, 1), "27-10-2026")
    set_cell_text(doc.tables[2].cell(0, 1), "Pendiente de aprobación")
    set_cell_text(doc.tables[2].cell(1, 1), "Pendiente")
    for c, value in enumerate(["27-09-2026", "1.0", "Versión inicial del DAS de Haircut Home Studio", "Equipo de Proyecto"]):
        set_cell_text(doc.tables[3].cell(1, c), value, size=8.5)

    # Reconstruir desde la tabla de contenidos.
    toc_source = next(p for p in doc.paragraphs if p.text.strip() == "Tabla de Contenidos")
    remove_from_paragraph(doc, toc_source)
    doc.add_page_break()
    add_heading(doc, "Tabla de Contenidos", 1)
    toc_p = doc.add_paragraph()
    add_toc(toc_p)
    doc.add_page_break()

    add_heading(doc, "1 Introducción", 1)
    add_heading(doc, "1.1 Contexto del problema", 2)
    add_body(doc, "Haircut Home Studio es una aplicación web orientada a la gestión de servicios de peluquería, reservas y atención en salón o a domicilio. El sistema centraliza la publicación del catálogo, la disponibilidad del peluquero, las solicitudes de clientes y el seguimiento de cada cita. Además, incorpora un simulador privado de cortes y coloración basado en inteligencia artificial, que permite generar una referencia visual antes de reservar.")
    add_body(doc, "La solución reemplaza un desarrollo PHP anterior por una aplicación Laravel 12 con responsabilidades separadas, validación en servidor, control de acceso por roles y persistencia relacional. La arquitectura debe mantener la consistencia de los horarios, proteger las fotografías y permitir que las funciones administrativas evolucionen sin afectar el proceso principal de reserva.")

    add_heading(doc, "1.2 Propósito", 2)
    add_body(doc, "Este documento describe la arquitectura de Haircut Home Studio mediante el modelo 4+1. Presenta las decisiones estructurales, los componentes principales, los procesos relevantes, la distribución física, los casos de uso significativos y los atributos de calidad que orientan la implementación.")

    add_heading(doc, "1.3 Ámbito", 2)
    add_body(doc, "El alcance comprende la aplicación web para clientes, administradores y peluqueros; el catálogo de servicios y productos; el asistente de reserva; la configuración de disponibilidad; la gestión de solicitudes y agenda; la exportación de reportes; y el simulador de imágenes con IA. El pago se realiza presencialmente y no forma parte de la arquitectura actual.")

    add_heading(doc, "1.4 Definiciones, acrónimos y abreviaciones", 2)
    add_table(doc, ["ACRÓNIMO", "DESCRIPCIÓN"], [
        ("DAS", "Documento de Arquitectura de Sistema."),
        ("IA", "Inteligencia artificial utilizada para generar una referencia visual de corte o color."),
        ("API", "Interfaz de programación utilizada para comunicarse con servicios externos."),
        ("ORM", "Mapeo objeto-relacional; Laravel Eloquent implementa el acceso a los datos."),
        ("SMTP", "Protocolo utilizado para el envío de correo electrónico."),
        ("MVC", "Patrón Modelo Vista Controlador utilizado por Laravel."),
    ], widths=[1.4, 5.0])

    add_heading(doc, "1.5 Referencias", 2)
    add_body(doc, "Los diagramas arquitectónicos incluidos fueron elaborados a partir del código fuente y de las reglas implementadas en Haircut Home Studio.")
    add_bullets(doc, [
        "Documentación del proyecto y archivo README.",
        "Rutas HTTP definidas en routes/web.php.",
        "Modelos, controladores, servicios, trabajos en cola y migraciones Laravel.",
        "Diagramas de las vistas lógica, de procesos, de despliegue, física y de casos de uso.",
    ])

    add_heading(doc, "1.6 Resumen ejecutivo", 2)
    add_body(doc, "La solución adopta una arquitectura web modular basada en Laravel. La interfaz Blade y JavaScript consume rutas controladas por middleware de autenticación y rol. Los controladores coordinan servicios de negocio para disponibilidad, reservas, cuotas y procesamiento de imágenes. Los modelos Eloquent persisten la información en MySQL o MariaDB. La integración con Gemini se encapsula mediante un contrato de proveedor y puede ejecutarse en cola, mientras que las fotografías se almacenan separando archivos públicos y privados.")

    add_heading(doc, "1.7 Representación", 2)
    add_bullets(doc, [
        ("Vista lógica: ", "describe las entidades del dominio y las interacciones de una solicitud de reserva mediante diagramas de clases, secuencia y comunicaciones."),
        ("Vista de procesos: ", "representa el flujo de actividades entre cliente, sistema y peluquero."),
        ("Vista de despliegue: ", "según la nomenclatura de la asignatura, organiza paquetes y componentes de software."),
        ("Vista física: ", "muestra los nodos de ejecución, almacenamiento y servicios externos."),
        ("Escenarios 4+1: ", "presenta los casos de uso relevantes para validar la arquitectura."),
    ])

    doc.add_page_break()
    add_heading(doc, "2 Metas y Restricciones de la Arquitectura", 1)
    add_heading(doc, "2.1 Metas de la arquitectura", 2)
    add_bullets(doc, [
        ("Consistencia: ", "evitar que dos solicitudes ocupen el mismo bloque mediante revalidación transaccional de la disponibilidad."),
        ("Seguridad: ", "restringir funciones por rol, validar las entradas y proteger credenciales, sesiones y fotografías."),
        ("Usabilidad: ", "guiar al cliente en una reserva de cinco pasos con información clara de servicio, fecha, referencia, lugar y confirmación."),
        ("Mantenibilidad: ", "separar controladores, servicios, modelos, proveedores e infraestructura para facilitar cambios y pruebas."),
        ("Disponibilidad funcional: ", "mantener reservas, catálogo y administración operativos aunque el proveedor de IA o el correo no estén disponibles."),
        ("Privacidad: ", "almacenar las imágenes de IA fuera del directorio público y eliminarlas de acuerdo con la política de retención."),
    ])

    add_heading(doc, "2.2 Restricciones de la arquitectura", 2)
    add_bullets(doc, [
        ("Plataforma: ", "PHP 8.2 o superior y Laravel 12."),
        ("Interfaz: ", "vistas Blade, CSS y JavaScript compilados mediante Vite."),
        ("Base de datos: ", "MySQL o MariaDB en el entorno de operación; SQLite en memoria para pruebas automatizadas."),
        ("Autenticación: ", "sesiones Laravel y middleware de roles cliente, administrador y peluquero."),
        ("Procesamiento de IA: ", "requiere una clave de Gemini para operación real, almacenamiento privado y un worker cuando el procesamiento es asíncrono."),
        ("Correo: ", "la notificación depende de un servidor SMTP configurado, pero una falla de correo no revierte el cambio de estado de la reserva."),
        ("Archivos: ", "el servidor debe disponer de almacenamiento persistente para referencias y generaciones de IA."),
    ])

    add_heading(doc, "2.3 Otros antecedentes y consideraciones", 2)
    add_body(doc, "La disponibilidad se calcula considerando días de atención, días bloqueados, meses habilitados, duración del servicio, horario de cierre y reservas existentes. Los servicios de color de larga duración admiten una regla controlada de solapamiento con un corte corto durante el último bloque de pose.")
    add_body(doc, "El simulador permite servicios activos de las categorías corte y color. Antes de procesar una fotografía, el sistema valida el archivo, elimina metadatos, comprueba el consentimiento y aplica cuotas por usuario, IP y consumo global. Las imágenes vencen automáticamente, salvo aquellas vinculadas temporalmente a una reserva.")

    doc.add_page_break()
    add_heading(doc, "3 Vista de Casos de Uso y Escenarios de Calidad", 1)
    add_heading(doc, "3.1 Modelo de casos de uso", 2)
    add_body(doc, "Los actores principales son Visitante, Cliente, Administrador y Peluquero. Como actores externos participan la API de Gemini y el servidor de correo. Los casos de uso se separan en autenticación, reserva, simulación de IA, operación del peluquero y administración del catálogo.")
    add_heading(doc, "3.2 Casos de uso relevantes", 2)
    add_table(doc, ["Código", "Nombre", "Actores", "Prioridad"], [
        ("CU-001", "Autenticación", "Visitante y usuario autenticado", "Alta"),
        ("CU-002", "Gestionar reservas", "Cliente", "Muy alta"),
        ("CU-003", "Generar simulación de IA", "Cliente y API de Gemini", "Alta"),
        ("CU-004", "Gestionar agenda y solicitudes", "Peluquero", "Muy alta"),
        ("CU-005", "Administrar catálogo", "Administrador", "Alta"),
    ], widths=[0.8, 2.1, 2.7, 0.9])

    add_heading(doc, "3.3 Escenarios de calidad relevantes", 2)
    quality_rows = [
        ("Seguridad", "Un usuario intenta acceder a una ruta de otro rol.", "El middleware bloquea el acceso y no expone datos del módulo."),
        ("Consistencia", "Dos clientes intentan reservar el mismo horario.", "La transacción bloquea las reservas del día; solo una solicitud se registra."),
        ("Privacidad", "Un usuario intenta abrir una imagen de IA ajena.", "El sistema valida la propiedad y responde sin revelar el archivo."),
        ("Disponibilidad", "Gemini o SMTP no responde.", "La falla se registra y no interrumpe las funciones centrales de reserva y administración."),
        ("Modificabilidad", "Se reemplaza el proveedor de imágenes.", "La implementación cambia detrás del contrato AiImageProvider sin modificar el flujo del controlador."),
        ("Usabilidad", "El cliente solicita una hora.", "El asistente presenta las decisiones en pasos y confirma el estado pendiente al finalizar."),
    ]
    add_table(doc, ["Atributo", "Estímulo", "Respuesta esperada"], quality_rows, widths=[1.2, 2.6, 3.1])

    doc.add_page_break()
    add_heading(doc, "4 Vista Lógica", 1)
    add_body(doc, "La vista lógica presenta las clases del dominio y la colaboración necesaria para crear una solicitud de reserva. La simulación de IA aparece como una referencia opcional previamente generada.")
    add_figure(doc, "Diagrama de Clases", DIAGRAMS / "Vista logica" / "Diagrama de Clases.png", width=6.15)
    add_figure(doc, "Diagrama de Secuencia", DIAGRAMS / "Vista logica" / "Diagrama de Secuencia.png", width=6.15, page_break=True)
    add_figure(doc, "Diagrama de Comunicaciones", DIAGRAMS / "Vista logica" / "Diagrama de Comunicaciones.png", width=6.15, page_break=True)

    doc.add_page_break()
    add_heading(doc, "5 Vista de Procesos", 1)
    add_body(doc, "El proceso principal coordina la autenticación, la elección del servicio, el cálculo y la revalidación del horario, el registro de una solicitud pendiente y la decisión posterior del peluquero. La generación completa de IA se mantiene como un proceso independiente.")
    add_figure(doc, "Diagrama de Actividades", DIAGRAMS / "Vista procesos" / "Diagrama de Actividades.png", width=5.9)

    doc.add_page_break()
    add_heading(doc, "6 Vista de Despliegue", 1)
    add_body(doc, "De acuerdo con la nomenclatura utilizada en la asignatura, esta vista describe la organización del código y los componentes que cooperan en la solución.")
    add_figure(doc, "Diagrama de Paquetes", DIAGRAMS / "Vista de despliegue" / "Diagrama de Paquetes.png", width=6.15)
    add_figure(doc, "Diagrama de Componentes", DIAGRAMS / "Vista de despliegue" / "Diagrama de Componentes.png", width=6.15, page_break=True)

    doc.add_page_break()
    add_heading(doc, "7 Vista Física", 1)
    add_body(doc, "La vista física separa los dispositivos cliente, el servidor web y de aplicación, el servidor de datos, el almacenamiento y los servicios externos. La distribución puede consolidarse en un mismo host durante desarrollo o separarse en producción.")
    add_figure(doc, "Diagrama de Despliegue", DIAGRAMS / "Vista fisica" / "Diagrama de Despliegue.png", width=6.15)

    doc.add_page_break()
    add_heading(doc, "8 Vista Escenarios 4+1", 1)
    add_body(doc, "Los siguientes diagramas verifican la arquitectura desde la perspectiva de los actores y sus objetivos.")
    use_cases = [
        ("CU-001 Autenticación", "Diagrama de Casos de Uso - Autenticación.png"),
        ("CU-002 Gestión de reservas", "Diagrama de Casos de Uso - Reservas.png"),
        ("CU-003 Simulador de IA", "Diagrama de Casos de Uso - Simulador IA.png"),
        ("CU-004 Operaciones del peluquero", "Diagrama de Casos de Uso - Peluquero.png"),
        ("CU-005 Administración del catálogo", "Diagrama de Casos de Uso - Administrador.png"),
    ]
    for i, (title, filename) in enumerate(use_cases):
        add_figure(doc, title, DIAGRAMS / "Casos de uso" / filename, width=6.0, page_break=(i > 0))

    doc.add_page_break()
    add_heading(doc, "9 Decisiones de Diseño y Selección de Alternativas", 1)
    add_heading(doc, "9.1 Framework web Laravel", 2)
    add_body(doc, "Se mantiene Laravel 12 porque proporciona una base consistente para enrutamiento, validación, sesiones, middleware, almacenamiento, colas, correo y pruebas. Esta decisión reduce código transversal y permite concentrar la implementación en las reglas del negocio.")
    add_heading(doc, "9.2 Separación de responsabilidades", 2)
    add_body(doc, "Los controladores reciben solicitudes HTTP y delegan las reglas complejas en servicios especializados. AvailabilityService concentra el cálculo de bloques y solapes; ReservationWorkflowService administra transiciones temporales; y los servicios de IA encapsulan cuotas, prompts, sanitización y proveedor externo.")
    add_heading(doc, "9.3 Consistencia de reservas", 2)
    add_body(doc, "La creación de una reserva revalida el horario dentro de una transacción y bloquea las reservas del día antes de insertar. Esta alternativa se eligió sobre una validación exclusiva en la interfaz porque protege frente a solicitudes concurrentes y cambios ocurridos después de mostrar el calendario.")
    add_heading(doc, "9.4 Procesamiento asíncrono de IA", 2)
    add_body(doc, "La generación se modela como un trabajo en cola. En demostración puede ejecutarse de forma síncrona, pero el modo asíncrono evita mantener la petición web abierta durante la comunicación con Gemini y permite aislar errores del proveedor.")
    add_heading(doc, "9.5 Proveedor intercambiable", 2)
    add_body(doc, "La integración de imágenes implementa el contrato AiImageProvider. Esta decisión permite utilizar un proveedor falso durante pruebas y sustituir Gemini sin acoplar el controlador o el trabajo en cola a una API concreta.")
    add_heading(doc, "9.6 Privacidad y retención", 2)
    add_body(doc, "Las imágenes del simulador se almacenan en un disco privado y se sirven únicamente después de comprobar al propietario o al personal autorizado. Un proceso programado elimina los archivos vencidos, reduciendo la exposición de datos personales y el consumo de almacenamiento.")
    add_heading(doc, "9.7 Dependencias externas no críticas", 2)
    add_body(doc, "El cambio de estado de una reserva no depende del éxito del correo. Del mismo modo, una falla en Gemini afecta la generación solicitada, pero no impide usar el catálogo, reservar, administrar solicitudes o consultar la agenda.")

    # Solicitar actualización de campos al abrir en Word.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
